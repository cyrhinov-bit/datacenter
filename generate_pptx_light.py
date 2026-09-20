import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_light_presentation():
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors - Light Institution Palette
    COLOR_BG = RGBColor(248, 250, 252)       # #F8FAFC Pure Slate Light
    COLOR_CARD = RGBColor(255, 255, 255)     # #FFFFFF Pure White
    COLOR_PRIMARY = RGBColor(15, 44, 89)     # #0F2C59 Deep Navy Text/Headers
    COLOR_CYAN = RGBColor(2, 132, 199)       # #0284C7 Azure / Tech Blue
    COLOR_EMERALD = RGBColor(5, 150, 105)    # #059669 Health Emerald
    COLOR_TEXT = RGBColor(51, 65, 85)        # #334155 Slate 700 Dark Text
    COLOR_MUTED = RGBColor(100, 116, 139)    # #64748B Slate 500 Subtext
    COLOR_BORDER = RGBColor(226, 232, 240)   # #E2E8F0 Soft Border
    COLOR_HEADER_BG = RGBColor(255, 255, 255)# White Header Bar

    logo_sante_path = os.path.abspath("assets/logo_sante.png")
    logo_disd_path = os.path.abspath("assets/logo_disd.png")

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, slide_num, slide_title, category="PROJET STRATÉGIQUE"):
        set_slide_background(slide)

        # Header Bar Container (Light)
        hdr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.1))
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = COLOR_HEADER_BG
        hdr.line.color.rgb = COLOR_BORDER

        # Add logos if present
        if os.path.exists(logo_sante_path):
            slide.shapes.add_picture(logo_sante_path, Inches(0.4), Inches(0.15), height=Inches(0.8))
        if os.path.exists(logo_disd_path):
            slide.shapes.add_picture(logo_disd_path, Inches(11.8), Inches(0.15), height=Inches(0.8))

        # Header Title in center
        txBox = slide.shapes.add_textbox(Inches(2.5), Inches(0.1), Inches(8.333), Inches(0.9))
        tf = txBox.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = f"{slide_num:02d} | {category.upper()}"
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_CYAN
        p1.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = slide_title.upper()
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_PRIMARY
        p2.alignment = PP_ALIGN.CENTER

        # Footer Bar
        ftr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.9), Inches(13.333), Inches(0.6))
        ftr.fill.solid()
        ftr.fill.fore_color.rgb = COLOR_HEADER_BG
        ftr.line.color.rgb = COLOR_BORDER

        ftr_tx = slide.shapes.add_textbox(Inches(0.5), Inches(6.95), Inches(12.333), Inches(0.5))
        ftr_tf = ftr_tx.text_frame
        p_ftr = ftr_tf.paragraphs[0]
        p_ftr.text = "Ministère de la Santé, de l'Hygiène Publique et de la CMU — Direction de l'Informatique et de la Santé Digitale (DISD)"
        p_ftr.font.size = Pt(10)
        p_ftr.font.color.rgb = COLOR_MUTED
        p_ftr.alignment = PP_ALIGN.CENTER

    def add_card(slide, left, top, width, height, title, subtitle, icon="🔹", accent_color=COLOR_CYAN):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1.5)

        txBox = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.15), width - Inches(0.3), height - Inches(0.3))
        tf = txBox.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"{icon}  {title}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = accent_color
        
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT
        p2.space_before = Pt(8)
        return card

    # ==================== SLIDE 1: Title / Garde ====================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    
    if os.path.exists(logo_sante_path):
        s1.shapes.add_picture(logo_sante_path, Inches(0.8), Inches(0.6), height=Inches(1.1))
    if os.path.exists(logo_disd_path):
        s1.shapes.add_picture(logo_disd_path, Inches(10.8), Inches(0.6), height=Inches(1.1))

    # Title Card Left
    tx1 = s1.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(7.5), Inches(4.5))
    tf1 = tx1.text_frame
    tf1.word_wrap = True
    
    p_badge = tf1.paragraphs[0]
    p_badge.text = "RÉPUBLIQUE DE CÔTE D'IVOIRE — MINISTÈRE DE LA SANTÉ"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_CYAN
    
    p_title = tf1.add_paragraph()
    p_title.text = "PROCESSUS DE DIGITALISATION"
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_PRIMARY
    p_title.space_before = Pt(12)

    p_sub = tf1.add_paragraph()
    p_sub.text = "du système de santé en Côte d'Ivoire"
    p_sub.font.size = Pt(22)
    p_sub.font.color.rgb = COLOR_EMERALD
    p_sub.space_before = Pt(4)

    p_pres = tf1.add_paragraph()
    p_pres.text = "\nPrésenté par : M. OUATTARA Yacouba\nIngénieur Informaticien\nDirection de l'Informatique et de la Santé Digitale (DISD)"
    p_pres.font.size = Pt(13)
    p_pres.font.color.rgb = COLOR_TEXT
    p_pres.space_before = Pt(16)

    # Visual Right Box
    card_r = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.6), Inches(2.0), Inches(3.9), Inches(4.5))
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = COLOR_CARD
    card_r.line.color.rgb = RGBColor(186, 230, 253)
    card_r.line.width = Pt(2)

    tx_r = s1.shapes.add_textbox(Inches(8.8), Inches(2.3), Inches(3.5), Inches(4.0))
    tf_r = tx_r.text_frame
    tf_r.word_wrap = True
    p_rt = tf_r.paragraphs[0]
    p_rt.text = "🏥  TRANSFORMATION DIGITALE"
    p_rt.font.size = Pt(16)
    p_rt.font.bold = True
    p_rt.font.color.rgb = COLOR_CYAN
    p_rt.alignment = PP_ALIGN.CENTER

    p_rc = tf_r.add_paragraph()
    p_rc.text = "\n• Datacenter National Haute Disponibilité\n\n• Interopérabilité des Établissements Hospitaliers\n\n• Centralisation et Sécurité des Données Sanitaires\n\n• Outils Modernes d'Aide à la Prise de Décision"
    p_rc.font.size = Pt(12)
    p_rc.font.color.rgb = COLOR_TEXT

    # ==================== SLIDE 2: Introduction ====================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, 2, "Introduction & Piliers Stratégiques", "Contexte Stratégique")

    b2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.333), Inches(1.0))
    b2.fill.solid()
    b2.fill.fore_color.rgb = COLOR_CARD
    b2.line.color.rgb = COLOR_CYAN
    b2_tx = s2.shapes.add_textbox(Inches(1.6), Inches(1.55), Inches(10.133), Inches(0.9))
    b2_tf = b2_tx.text_frame
    b2_tf.word_wrap = True
    p_b2 = b2_tf.paragraphs[0]
    p_b2.text = "Depuis 2021, le MSHPCMU accélère l'intégration du numérique dans le Système National d'Information Sanitaire (SNIS)."
    p_b2.font.size = Pt(15)
    p_b2.font.bold = True
    p_b2.font.color.rgb = COLOR_PRIMARY
    p_b2.alignment = PP_ALIGN.CENTER

    w4 = Inches(2.7)
    add_card(s2, Inches(0.8), Inches(2.8), w4, Inches(3.6), "Qualité des données", "Améliorer la qualité, la fiabilité et la promptitude des statistiques sanitaires collectées.", "💾", COLOR_CYAN)
    add_card(s2, Inches(3.8), Inches(2.8), w4, Inches(3.6), "Prise de décision", "Faciliter le pilotage stratégique et opérationnel basé sur des indicateurs en temps réel.", "📊", RGBColor(126, 34, 206))
    add_card(s2, Inches(6.8), Inches(2.8), w4, Inches(3.6), "Prise en charge", "Fluidifier le parcours patient et optimiser la continuité des soins dans les établissements.", "🏥", COLOR_EMERALD)
    add_card(s2, Inches(9.8), Inches(2.8), w4, Inches(3.6), "Modernisation", "Moderniser et pérenniser globalement le système d'information de santé national.", "💻", COLOR_CYAN)

    # ==================== SLIDE 3: Le Datacenter ====================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, 3, "Le Datacenter — Cœur Technique", "Infrastructure")

    center_hub = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.9), Inches(2.7), Inches(3.5), Inches(2.2))
    center_hub.fill.solid()
    center_hub.fill.fore_color.rgb = COLOR_CARD
    center_hub.line.color.rgb = COLOR_CYAN
    center_hub.line.width = Pt(3)

    c_tx = s3.shapes.add_textbox(Inches(5.0), Inches(2.9), Inches(3.3), Inches(1.8))
    c_tf = c_tx.text_frame
    c_tf.word_wrap = True
    p_c1 = c_tf.paragraphs[0]
    p_c1.text = "🖥️  DATACENTER"
    p_c1.font.size = Pt(18)
    p_c1.font.bold = True
    p_c1.font.color.rgb = COLOR_CYAN
    p_c1.alignment = PP_ALIGN.CENTER
    p_c2 = c_tf.add_paragraph()
    p_c2.text = "Noyau de Gestion des Données de Santé"
    p_c2.font.size = Pt(11)
    p_c2.font.color.rgb = COLOR_PRIMARY
    p_c2.alignment = PP_ALIGN.CENTER

    add_card(s3, Inches(4.9), Inches(1.35), Inches(3.5), Inches(1.15), "Centraliser les données", "Entrepôt unifié national", "💾", COLOR_CYAN)
    add_card(s3, Inches(0.8), Inches(3.2), Inches(3.6), Inches(1.3), "Sécuriser les informations", "Confidentialité, résilience et PSSI", "🔒", COLOR_EMERALD)
    add_card(s3, Inches(8.9), Inches(3.2), Inches(3.6), Inches(1.3), "Déployer des systèmes intégrés", "Interopérabilité et communication", "⚙️", RGBColor(126, 34, 206))
    add_card(s3, Inches(4.9), Inches(5.15), Inches(3.5), Inches(1.2), "Moderniser les outils", "Outils de gestion performants", "🖥️", COLOR_CYAN)

    # ==================== SLIDE 4: Politique de Digitalisation ====================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, 4, "Politique de Digitalisation", "Vision Sectorielle")

    add_card(s4, Inches(0.8), Inches(1.4), Inches(3.6), Inches(1.2), "Administration", "Pilotage RH, finances et gouvernance", "🏢", COLOR_CYAN)
    add_card(s4, Inches(4.8), Inches(1.4), Inches(3.7), Inches(1.2), "DIGITALISATION GLOBALE", "Convergence des deux versants", "🌐", COLOR_EMERALD)
    add_card(s4, Inches(8.9), Inches(1.4), Inches(3.6), Inches(1.2), "Offre de soins", "Hôpitaux, cliniques et consultations", "🏥", COLOR_CYAN)

    add_card(s4, Inches(0.8), Inches(2.9), Inches(2.7), Inches(3.5), "Projets & Apps", "Déploiement des logiciels hospitaliers et applicatifs métiers.", "💻", COLOR_CYAN)
    add_card(s4, Inches(3.8), Inches(2.9), Inches(2.7), Inches(3.5), "Sécurité des SI", "Application rigoureuse de la PSSI et audits de cyberdéfense.", "🔐", COLOR_EMERALD)
    add_card(s4, Inches(6.8), Inches(2.9), Inches(2.7), Inches(3.5), "Archivage Numérique", "Gestion électronique des documents (GED) et traçabilité.", "📁", RGBColor(126, 34, 206))
    add_card(s4, Inches(9.8), Inches(2.9), Inches(2.7), Inches(3.5), "Données Centralisées", "Consolidation nationale pour un pilotage sanitaire continu.", "☁️", COLOR_CYAN)

    # ==================== SLIDE 5: Transformation du SNIS ====================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, 5, "Transformation du SNIS", "Écosystème Applicatif")

    b5 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.4), Inches(10.333), Inches(0.9))
    b5.fill.solid()
    b5.fill.fore_color.rgb = COLOR_CARD
    b5.line.color.rgb = COLOR_CYAN
    b5_tx = s5.shapes.add_textbox(Inches(1.6), Inches(1.45), Inches(10.133), Inches(0.8))
    p_b5 = b5_tx.text_frame.paragraphs[0]
    p_b5.text = "🎯 Le SNIS évolue pour rendre les informations sanitaires : DISPONIBLES • FIABLES • ACCESSIBLES À TEMPS"
    p_b5.font.size = Pt(13)
    p_b5.font.bold = True
    p_b5.font.color.rgb = COLOR_PRIMARY
    p_b5.alignment = PP_ALIGN.CENTER

    w3 = Inches(3.6)
    add_card(s5, Inches(0.8), Inches(2.6), w3, Inches(3.8), "DHIS2", "Collecte globale, agrégation et analyse des données sanitaires et indicateurs nationaux.", "📊", COLOR_CYAN)
    add_card(s5, Inches(4.8), Inches(2.6), w3, Inches(3.8), "SIGDEP3", "Système intégré pour le suivi clinique et biologique individualisé des patients (VIH).", "👤", COLOR_EMERALD)
    add_card(s5, Inches(8.9), Inches(2.6), w3, Inches(3.8), "OPENELIS", "Gestion informatisée des laboratoires de biologie médicale et du circuit des échantillons.", "🧪", RGBColor(126, 34, 206))

    # ==================== SLIDE 6: Autres Solutions Numériques ====================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, 6, "Autres Solutions Numériques", "Complémentarité Métier")

    w2 = Inches(5.6)
    h2 = Inches(1.8)
    add_card(s6, Inches(0.8), Inches(1.6), w2, h2, "E-SIGL", "Gestion logistique intégrée des pharmacies et approvisionnement en intrants stratégiques.", "💊", COLOR_CYAN)
    add_card(s6, Inches(6.9), Inches(1.6), w2, h2, "M-SUPPLY", "Gestion des stocks et traçabilité de la dispensation des médicaments au comptoir.", "💉", COLOR_EMERALD)
    add_card(s6, Inches(0.8), Inches(3.7), w2, h2, "UPID", "Identifiant Unique du Patient (notamment sous ARV) pour l'unicité et le suivi longitudinal.", "👤", RGBColor(126, 34, 206))
    add_card(s6, Inches(6.9), Inches(3.7), w2, h2, "DPI", "Dossier Patient Informatisé pour la continuité des soins et l'historique médical partagé.", "📋", COLOR_CYAN)

    b6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.5), Inches(5.8), Inches(8.333), Inches(0.7))
    b6.fill.solid()
    b6.fill.fore_color.rgb = COLOR_CARD
    b6.line.color.rgb = COLOR_EMERALD
    p_b6 = b6.text_frame.paragraphs[0]
    p_b6.text = "💡 Des solutions complémentaires au service de la transformation numérique."
    p_b6.font.size = Pt(12)
    p_b6.font.bold = True
    p_b6.font.color.rgb = COLOR_PRIMARY
    p_b6.alignment = PP_ALIGN.CENTER

    # ==================== SLIDE 7: Mise en Œuvre du Datacenter ====================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, 7, "Mise en Œuvre du Datacenter", "Socle d'Infrastructure")

    b7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.333), Inches(1.3))
    b7.fill.solid()
    b7.fill.fore_color.rgb = COLOR_CARD
    b7.line.color.rgb = COLOR_CYAN
    b7_tx = s7.shapes.add_textbox(Inches(1.6), Inches(1.6), Inches(10.133), Inches(1.1))
    p_b7 = b7_tx.text_frame.paragraphs[0]
    p_b7.text = "🖥️  INFRASTRUCTURE NATIONALE DE SANTÉ"
    p_b7.font.size = Pt(16)
    p_b7.font.bold = True
    p_b7.font.color.rgb = COLOR_CYAN
    p_b7.alignment = PP_ALIGN.CENTER
    p_b7_sub = b7_tx.text_frame.add_paragraph()
    p_b7_sub.text = "« Mettre en place une infrastructure centralisée, sécurisée et performante pour accompagner la transformation numérique du système de santé. »"
    p_b7_sub.font.size = Pt(12)
    p_b7_sub.font.color.rgb = COLOR_PRIMARY
    p_b7_sub.alignment = PP_ALIGN.CENTER

    add_card(s7, Inches(0.8), Inches(3.2), w3, Inches(3.2), "CENTRALISÉ", "Consolidation des données et mutualisation des ressources serveur.", "💾", COLOR_CYAN)
    add_card(s7, Inches(4.8), Inches(3.2), w3, Inches(3.2), "SÉCURISÉ", "Protection maximale, chiffrement et haute résilience contre les sinistres.", "🔒", COLOR_EMERALD)
    add_card(s7, Inches(8.9), Inches(3.2), w3, Inches(3.2), "PERFORMANT", "Temps de réponse optimaux et disponibilité continue 24/7.", "⚡", COLOR_CYAN)

    # ==================== SLIDE 8: Chronologie de mise en œuvre ====================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, 8, "Chronologie de Mise en Œuvre (2021 - 2024)", "Feuille de Route")

    add_card(s8, Inches(0.8), Inches(1.8), w4, Inches(4.2), "2021 — Cadrage", "Signature des contrats, cadrage institutionnel et passation des marchés.\n\nPhase : Préparation", "📄", COLOR_CYAN)
    add_card(s8, Inches(3.8), Inches(1.8), w4, Inches(4.2), "2022 — Réalisation", "Démarrage des travaux d'aménagement, génie civil et pose des équipements.\n\nPhase : Travaux & Équipements", "🏗️", RGBColor(217, 119, 6))
    add_card(s8, Inches(6.8), Inches(1.8), w4, Inches(4.2), "2023 — Recette", "Tests de recette technique en usine, tests de charge et validation sécuritaire.\n\nPhase : Tests de Recette", "🧪", RGBColor(126, 34, 206))
    add_card(s8, Inches(9.8), Inches(1.8), w4, Inches(4.2), "2024 — Service", "Réception provisoire des installations et mise en service opérationnelle.\n\nPhase : Exploitation Active", "✅", COLOR_EMERALD)

    # ==================== SLIDE 9: Le Datacenter en Images ====================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, 9, "Le Datacenter en Images", "Galerie Visuelle")

    add_card(s9, Inches(0.8), Inches(1.8), w3, Inches(4.2), "Salle Serveurs & Baies IT", "Photo 1 : Racks haute densité, connectivité fibre et câblage structuré certifié.", "📸", COLOR_CYAN)
    add_card(s9, Inches(4.8), Inches(1.8), w3, Inches(4.2), "Énergie & Onduleurs", "Photo 2 : Alimentation électrique secourue, redondance N+1 et groupes électrogènes.", "⚡", RGBColor(217, 119, 6))
    add_card(s9, Inches(8.9), Inches(1.8), w3, Inches(4.2), "Supervision & Contrôle", "Photo 3 : Poste de monitoring 24/7 et contrôle d'accès biométrique strict.", "🛡️", COLOR_EMERALD)

    # ==================== SLIDE 10: Infrastructure technique ====================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, 10, "Infrastructure Technique du Datacenter", "Spécifications Équipements")

    add_card(s10, Inches(0.8), Inches(1.8), w3, Inches(4.0), "Puissance de Calcul", "Serveurs de dernière génération pour exécuter les charges critiques du SNIS.", "💻", COLOR_CYAN)
    add_card(s10, Inches(4.8), Inches(1.8), w3, Inches(4.0), "Refroidissement de Précision", "Climatisation redondante avec confinement des allées chaudes/froides.", "❄️", COLOR_CYAN)
    add_card(s10, Inches(8.9), Inches(1.8), w3, Inches(4.0), "Sécurité Incendie & Physique", "Extinction automatique par gaz inerte et détection ultra-précoce des fumées.", "🔒", COLOR_EMERALD)

    # ==================== SLIDE 11: Une infrastructure au service de la santé ====================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, 11, "Au Service de la Santé Digitale", "Valeur Ajoutée")

    add_card(s11, Inches(0.8), Inches(1.6), Inches(5.5), Inches(4.8), "SOCLE DE CONFIANCE", "Le Datacenter constitue une infrastructure essentielle pour accompagner le développement des systèmes d'information sanitaires et garantir un service continu aux établissements.", "🛡️", COLOR_CYAN)
    add_card(s11, Inches(6.8), Inches(1.6), Inches(5.7), Inches(1.4), "Centralisation", "Unification complète des bases de données sanitaires nationales.", "💾", COLOR_CYAN)
    add_card(s11, Inches(6.8), Inches(3.3), Inches(5.7), Inches(1.4), "Sécurisation", "Chiffrement et conformité stricte aux exigences de cybersécurité.", "🔒", COLOR_EMERALD)
    add_card(s11, Inches(6.8), Inches(5.0), Inches(5.7), Inches(1.4), "Performance", "Temps d'accès instantanés pour les professionnels de santé.", "⚡", COLOR_CYAN)

    # ==================== SLIDE 12: Défis et Perspectives ====================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, 12, "Défis et Perspectives", "Enjeux Stratégiques")

    w5 = Inches(2.2)
    add_card(s12, Inches(0.6), Inches(2.0), w5, Inches(4.2), "Ressources Humaines", "Recrutement et fidélisation d'experts IT qualifiés.", "👥", COLOR_CYAN)
    add_card(s12, Inches(3.0), Inches(2.0), w5, Inches(4.2), "Formation Continue", "Montée en compétences des agents et du personnel de santé.", "🎓", RGBColor(217, 119, 6))
    add_card(s12, Inches(5.4), Inches(2.0), w5, Inches(4.2), "Maintenance", "Maintenance préventive, SLA et gestion des pièces de rechange.", "🔧", RGBColor(234, 88, 12))
    add_card(s12, Inches(7.8), Inches(2.0), w5, Inches(4.2), "Interconnexion", "Raccordement réseau de tous les centres de santé périphériques.", "🔗", RGBColor(126, 34, 206))
    add_card(s12, Inches(10.2), Inches(2.0), w5, Inches(4.2), "Sécurité Données", "Conformité PSSI, audits réguliers et cyber-résilience continue.", "🔒", COLOR_EMERALD)

    # ==================== SLIDE 13: Connectivité et Interconnexion ====================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, 13, "Connectivité & Interconnexion", "Maillage Réseau")

    w6 = Inches(3.6)
    h6 = Inches(2.2)
    add_card(s13, Inches(0.8), Inches(1.6), w6, h6, "Fibre Optique", "Dorsale nationale haut débit pour les CHU, CHR et hôpitaux généraux.", "🌐", COLOR_CYAN)
    add_card(s13, Inches(4.8), Inches(1.6), w6, h6, "VSAT Gouvernemental", "Liaisons satellitaires pour les localités enclavées.", "🛰️", RGBColor(126, 34, 206))
    add_card(s13, Inches(8.9), Inches(1.6), w6, h6, "Internet Haut Débit", "Garantie de bande passante et latence minimale.", "📡", COLOR_CYAN)
    add_card(s13, Inches(0.8), Inches(4.1), w6, h6, "Couverture 4G", "Tablettes terrain et accès mobile pour la collecte décentralisée.", "📱", COLOR_EMERALD)
    add_card(s13, Inches(4.8), Inches(4.1), w6, h6, "Coût de la Data", "Tarification préférentielle et partenariats opérateurs.", "💰", RGBColor(217, 119, 6))
    add_card(s13, Inches(8.9), Inches(4.1), w6, h6, "Alimentation Électrique", "Stabilité du réseau électrique et solutions solaires d'appoint.", "⚡", RGBColor(234, 88, 12))

    # ==================== SLIDE 14: Identité, Adoption et Sécurité ====================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, 14, "Identité, Adoption et Sécurité", "Gouvernance & Confiance")

    add_card(s14, Inches(0.8), Inches(2.0), w4, Inches(4.2), "Identifiant Unique", "Garantir l'unicité du patient et éviter les doublons dans le système sanitaire.\n\n→ IDENTIFIER", "👤", RGBColor(126, 34, 206))
    add_card(s14, Inches(3.8), Inches(2.0), w4, Inches(4.2), "Appropriation Outils", "Conduite du changement active et accompagnement quotidien des personnels.\n\n→ ADOPTER", "👥", COLOR_CYAN)
    add_card(s14, Inches(6.8), Inches(2.0), w4, Inches(4.2), "Sécurité des Données", "Protection absolue du secret médical, contrôle d'accès strict et chiffrement.\n\n→ SÉCURISER", "🔒", COLOR_EMERALD)
    add_card(s14, Inches(9.8), Inches(2.0), w4, Inches(4.2), "Cadre Réglementaire", "Conformité légale sur la protection des données personnelles (ARTCI).\n\n→ ENCADRER", "⚖️", COLOR_CYAN)

    # ==================== SLIDE 15: Conclusion ====================
    s15 = prs.slides.add_slide(blank_layout)
    add_header(s15, 15, "Conclusion & Perspectives", "Clôture")

    b15 = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.333), Inches(1.6))
    b15.fill.solid()
    b15.fill.fore_color.rgb = COLOR_CARD
    b15.line.color.rgb = COLOR_CYAN
    b15_tx = s15.shapes.add_textbox(Inches(1.6), Inches(1.7), Inches(10.133), Inches(1.2))
    p15 = b15_tx.text_frame.paragraphs[0]
    p15.text = "« La digitalisation transforme progressivement le système de santé ivoirien. »"
    p15.font.size = Pt(22)
    p15.font.bold = True
    p15.font.color.rgb = COLOR_PRIMARY
    p15.alignment = PP_ALIGN.CENTER

    add_card(s15, Inches(0.8), Inches(3.4), w4, Inches(2.2), "Accès à l'Information", "Données sanitaires fiables en temps réel pour tous les acteurs.", "📖", COLOR_CYAN)
    add_card(s15, Inches(3.8), Inches(3.4), w4, Inches(2.2), "Qualité des Soins", "Prise en charge patient fluidifiée et suivi médical continu.", "🩺", COLOR_EMERALD)
    add_card(s15, Inches(6.8), Inches(3.4), w4, Inches(2.2), "Gestion du Système", "Efficience accrue des ressources matérielles et humaines.", "📊", RGBColor(217, 119, 6))
    add_card(s15, Inches(9.8), Inches(3.4), w4, Inches(2.2), "Prise de Décision", "Politiques publiques de santé fondées sur des preuves tangibles.", "🎯", RGBColor(126, 34, 206))

    cta = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(5.9), Inches(10.333), Inches(0.7))
    cta.fill.solid()
    cta.fill.fore_color.rgb = COLOR_CARD
    cta.line.color.rgb = COLOR_EMERALD
    p_cta = cta.text_frame.paragraphs[0]
    p_cta.text = "✨ La réussite du projet repose sur l'engagement et la contribution de tous les acteurs."
    p_cta.font.size = Pt(13)
    p_cta.font.bold = True
    p_cta.font.color.rgb = COLOR_EMERALD
    p_cta.alignment = PP_ALIGN.CENTER

    output_path = "Presentation_Datacenter_Sante_CI_Light.pptx"
    prs.save(output_path)
    print(f"Light presentation successfully saved to {output_path}")

if __name__ == "__main__":
    create_light_presentation()

