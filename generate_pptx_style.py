import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_styled_presentation():
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palette extracted from style.jpeg
    COLOR_BG = RGBColor(240, 243, 246)       # #F0F3F6 Soft Background
    COLOR_WHITE = RGBColor(255, 255, 255)    # #FFFFFF Pure White
    COLOR_RED = RGBColor(220, 20, 30)        # #DC141E Red
    COLOR_TEAL = RGBColor(0, 122, 135)       # #007A87 Teal
    COLOR_ORANGE = RGBColor(245, 80, 10)     # #F5500A Orange
    COLOR_GREEN = RGBColor(60, 165, 50)      # #3CA532 Green
    COLOR_DARK = RGBColor(16, 42, 56)        # #102A38 Dark Teal
    COLOR_BORDER = RGBColor(218, 225, 233)   # Soft Gray border
    COLOR_TEXT_DARK = RGBColor(20, 35, 45)   # Dark Charcoal text
    COLOR_MUTED = RGBColor(100, 116, 139)    # Muted slate text

    logo_sante_path = os.path.abspath("assets/logo_sante.png")
    logo_disd_path = os.path.abspath("assets/logo_disd.png")
    img_dc1 = os.path.abspath("assets/datacenter1.png")
    img_dc2 = os.path.abspath("assets/datacenter2.png")
    img_dc3 = os.path.abspath("assets/datacenter3.png")
    img_dc4 = os.path.abspath("assets/datacenter4.png")
    img_dc11 = os.path.abspath("assets/datacenter_slide_11_image_3.png")

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, slide_num, slide_title, category="PROJET STRATÉGIQUE"):
        set_slide_background(slide)

        hdr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.15))
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = COLOR_WHITE
        hdr.line.color.rgb = COLOR_BORDER

        if os.path.exists(logo_sante_path):
            slide.shapes.add_picture(logo_sante_path, Inches(0.4), Inches(0.18), height=Inches(0.78))
        if os.path.exists(logo_disd_path):
            slide.shapes.add_picture(logo_disd_path, Inches(11.8), Inches(0.18), height=Inches(0.78))

        txBox = slide.shapes.add_textbox(Inches(2.5), Inches(0.1), Inches(8.333), Inches(0.95))
        tf = txBox.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = f"{slide_num:02d}  •  {category.upper()}"
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEAL
        p1.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = slide_title.upper()
        p2.font.size = Pt(17)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_DARK
        p2.alignment = PP_ALIGN.CENTER

        ftr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.9), Inches(13.333), Inches(0.6))
        ftr.fill.solid()
        ftr.fill.fore_color.rgb = COLOR_WHITE
        ftr.line.color.rgb = COLOR_BORDER

        ftr_tx = slide.shapes.add_textbox(Inches(0.5), Inches(6.95), Inches(12.333), Inches(0.5))
        ftr_tf = ftr_tx.text_frame
        p_ftr = ftr_tf.paragraphs[0]
        p_ftr.text = "Ministère de la Santé, de l'Hygiène Publique et de la CMU  |  Direction de l'Informatique et de la Santé Digitale (DISD)"
        p_ftr.font.size = Pt(10)
        p_ftr.font.bold = True
        p_ftr.font.color.rgb = COLOR_MUTED
        p_ftr.alignment = PP_ALIGN.CENTER

    def add_horizontal_pill(slide, left, top, width, height, number_str, title, subtitle="", color=COLOR_RED):
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        pill.fill.solid()
        pill.fill.fore_color.rgb = color
        pill.line.fill.background()

        badge_diameter = min(height - Inches(0.18), Inches(0.72))
        badge_left = left + Inches(0.12)
        badge_top = top + (height - badge_diameter) / 2
        
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, badge_left, badge_top, badge_diameter, badge_diameter)
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_WHITE
        badge.line.fill.background()

        badge_tf = badge.text_frame
        badge_tf.margin_left = Inches(0)
        badge_tf.margin_right = Inches(0)
        badge_tf.margin_top = Inches(0)
        badge_tf.margin_bottom = Inches(0)
        badge_tf.word_wrap = False
        
        p_num = badge_tf.paragraphs[0]
        p_num.text = str(number_str)
        p_num.font.size = Pt(15 if len(str(number_str)) <= 2 else 11)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_DARK
        p_num.alignment = PP_ALIGN.CENTER

        tx_left = badge_left + badge_diameter + Inches(0.16)
        tx_width = width - (tx_left - left) - Inches(0.15)
        txBox = slide.shapes.add_textbox(tx_left, top + Inches(0.08), tx_width, height - Inches(0.16))
        tf = txBox.text_frame
        tf.margin_left = Inches(0)
        tf.margin_right = Inches(0)
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title.upper()
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE

        if subtitle:
            p_s = tf.add_paragraph()
            p_s.text = subtitle
            p_s.font.size = Pt(10)
            p_s.font.color.rgb = RGBColor(240, 245, 250)
            p_s.space_before = Pt(2)

        return pill

    def add_vertical_pill(slide, left, top, width, height, number_str, title, subtitle="", color=COLOR_RED, tag=""):
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        pill.fill.solid()
        pill.fill.fore_color.rgb = color
        pill.line.fill.background()

        badge_diameter = Inches(0.68)
        badge_left = left + (width - badge_diameter) / 2
        badge_top = top + Inches(0.15)

        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, badge_left, badge_top, badge_diameter, badge_diameter)
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_WHITE
        badge.line.fill.background()

        badge_tf = badge.text_frame
        badge_tf.margin_left = Inches(0)
        badge_tf.margin_right = Inches(0)
        badge_tf.margin_top = Inches(0)
        badge_tf.margin_bottom = Inches(0)
        badge_tf.word_wrap = False

        p_num = badge_tf.paragraphs[0]
        p_num.text = str(number_str)
        p_num.font.size = Pt(15 if len(str(number_str)) <= 2 else 11)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_DARK
        p_num.alignment = PP_ALIGN.CENTER

        tx_top = badge_top + badge_diameter + Inches(0.12)
        tx_height = height - (tx_top - top) - Inches(0.15)
        txBox = slide.shapes.add_textbox(left + Inches(0.12), tx_top, width - Inches(0.24), tx_height)
        tf = txBox.text_frame
        tf.margin_left = Inches(0)
        tf.margin_right = Inches(0)
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title.upper()
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE
        p_t.alignment = PP_ALIGN.CENTER

        if subtitle:
            p_s = tf.add_paragraph()
            p_s.text = subtitle
            p_s.font.size = Pt(9.5)
            p_s.font.color.rgb = RGBColor(240, 245, 250)
            p_s.space_before = Pt(4)
            p_s.alignment = PP_ALIGN.CENTER

        if tag:
            p_tag = tf.add_paragraph()
            p_tag.text = tag.upper()
            p_tag.font.size = Pt(8.5)
            p_tag.font.bold = True
            p_tag.font.color.rgb = COLOR_WHITE
            p_tag.space_before = Pt(6)
            p_tag.alignment = PP_ALIGN.CENTER

        return pill

    def add_image_card(slide, left, top, width, height, image_path, number_str, title, subtitle, color=COLOR_RED):
        # Container Card
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1.5)

        # Image inside top portion
        img_top = top + Inches(0.1)
        img_left = left + Inches(0.1)
        img_width = width - Inches(0.2)
        img_height = height * 0.58

        if os.path.exists(image_path):
            slide.shapes.add_picture(image_path, img_left, img_top, width=img_width, height=img_height)

        # Bottom Capsule Header inside Card
        pill_top = img_top + img_height + Inches(0.1)
        pill_height = height - (pill_top - top) - Inches(0.1)
        
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, img_left, pill_top, img_width, pill_height)
        pill.fill.solid()
        pill.fill.fore_color.rgb = color
        pill.line.fill.background()

        badge_diameter = min(pill_height - Inches(0.14), Inches(0.55))
        badge_left = img_left + Inches(0.08)
        badge_top = pill_top + (pill_height - badge_diameter) / 2

        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, badge_left, badge_top, badge_diameter, badge_diameter)
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_WHITE
        badge.line.fill.background()

        badge_tf = badge.text_frame
        badge_tf.margin_left = Inches(0)
        badge_tf.margin_right = Inches(0)
        badge_tf.margin_top = Inches(0)
        badge_tf.margin_bottom = Inches(0)
        p_num = badge_tf.paragraphs[0]
        p_num.text = str(number_str)
        p_num.font.size = Pt(12)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_DARK
        p_num.alignment = PP_ALIGN.CENTER

        tx_left = badge_left + badge_diameter + Inches(0.12)
        tx_width = img_width - (tx_left - img_left) - Inches(0.1)
        txBox = slide.shapes.add_textbox(tx_left, pill_top + Inches(0.04), tx_width, pill_height - Inches(0.08))
        tf = txBox.text_frame
        tf.margin_left = Inches(0)
        tf.margin_right = Inches(0)
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title.upper()
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE

        if subtitle:
            p_s = tf.add_paragraph()
            p_s.text = subtitle
            p_s.font.size = Pt(8.5)
            p_s.font.color.rgb = RGBColor(240, 245, 250)

    # ==================== SLIDE 1: Title / Page de Garde ====================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    if os.path.exists(logo_sante_path):
        s1.shapes.add_picture(logo_sante_path, Inches(0.8), Inches(0.6), height=Inches(1.1))
    if os.path.exists(logo_disd_path):
        s1.shapes.add_picture(logo_disd_path, Inches(10.8), Inches(0.6), height=Inches(1.1))

    tx1 = s1.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(7.5), Inches(4.5))
    tf1 = tx1.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "RÉPUBLIQUE DE CÔTE D'IVOIRE — MINISTÈRE DE LA SANTÉ"
    p_badge.font.size = Pt(11)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_TEAL

    p_title = tf1.add_paragraph()
    p_title.text = "PROCESSUS DE DIGITALISATION"
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_DARK
    p_title.space_before = Pt(10)

    p_sub = tf1.add_paragraph()
    p_sub.text = "du système de santé en Côte d'Ivoire"
    p_sub.font.size = Pt(22)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_RED
    p_sub.space_before = Pt(4)

    p_pres = tf1.add_paragraph()
    p_pres.text = "\nPrésenté par : M. OUATTARA Yacouba\nIngénieur Informaticien\nDirection de l'Informatique et de la Santé Digitale (DISD)"
    p_pres.font.size = Pt(13)
    p_pres.font.color.rgb = COLOR_TEXT_DARK
    p_pres.space_before = Pt(14)

    card_r = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.6), Inches(1.9), Inches(3.9), Inches(4.8))
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = COLOR_WHITE
    card_r.line.color.rgb = COLOR_BORDER
    card_r.line.width = Pt(1.5)

    tx_rt = s1.shapes.add_textbox(Inches(8.8), Inches(2.1), Inches(3.5), Inches(0.5))
    p_rt = tx_rt.text_frame.paragraphs[0]
    p_rt.text = "🏥  TRANSFORMATION NUMÉRIQUE"
    p_rt.font.size = Pt(12)
    p_rt.font.bold = True
    p_rt.font.color.rgb = COLOR_DARK
    p_rt.alignment = PP_ALIGN.CENTER

    add_horizontal_pill(s1, Inches(8.8), Inches(2.75), Inches(3.5), Inches(0.85), "1", "Datacenter National", "Haute disponibilité & Sécurité", COLOR_RED)
    add_horizontal_pill(s1, Inches(8.8), Inches(3.7), Inches(3.5), Inches(0.85), "2", "Systèmes Intégrés", "Interopérabilité SNIS & Hôpitaux", COLOR_TEAL)
    add_horizontal_pill(s1, Inches(8.8), Inches(4.65), Inches(3.5), Inches(0.85), "3", "Décision & Soins", "Pilotage par la donnée fiable", COLOR_GREEN)
    add_horizontal_pill(s1, Inches(8.8), Inches(5.6), Inches(3.5), Inches(0.85), "4", "Gouvernance & SSI", "Pérénnisation & Confiance", COLOR_DARK)

    # ==================== SLIDE 2: Introduction ====================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, 2, "Introduction & Piliers Stratégiques", "Contexte Stratégique")

    b2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.4), Inches(10.9), Inches(0.95))
    b2.fill.solid()
    b2.fill.fore_color.rgb = COLOR_WHITE
    b2.line.color.rgb = COLOR_TEAL
    b2.line.width = Pt(2)
    b2_tx = s2.shapes.add_textbox(Inches(1.3), Inches(1.45), Inches(10.7), Inches(0.85))
    p_b2 = b2_tx.text_frame.paragraphs[0]
    p_b2.text = "« Depuis 2021, le MSHPCMU accélère l'intégration du numérique dans le Système National d'Information Sanitaire (SNIS). »"
    p_b2.font.size = Pt(14)
    p_b2.font.bold = True
    p_b2.font.color.rgb = COLOR_DARK
    p_b2.alignment = PP_ALIGN.CENTER

    add_horizontal_pill(s2, Inches(1.2), Inches(2.6), Inches(5.3), Inches(1.8), "1", "Qualité des données", "Améliorer l'exhaustivité, la promptitude et la fiabilité des statistiques sanitaires.", COLOR_RED)
    add_horizontal_pill(s2, Inches(6.8), Inches(2.6), Inches(5.3), Inches(1.8), "2", "Prise de décision", "Faciliter le pilotage stratégique et opérationnel basé sur des indicateurs fiables.", COLOR_TEAL)
    add_horizontal_pill(s2, Inches(1.2), Inches(4.7), Inches(5.3), Inches(1.8), "3", "Prise en charge des soins", "Fluidifier le parcours patient et optimiser la continuité des soins dans les centres.", COLOR_ORANGE)
    add_horizontal_pill(s2, Inches(6.8), Inches(4.7), Inches(5.3), Inches(1.8), "4", "Modernisation du système", "Moderniser et pérenniser durablement l'ensemble du système de santé ivoirien.", COLOR_GREEN)

    # ==================== SLIDE 3: Le Datacenter ====================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, 3, "Le Datacenter — Hub Central", "Cœur Technique")

    hub_circle = s3.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), Inches(2.2), Inches(3.4), Inches(3.4))
    hub_circle.fill.solid()
    hub_circle.fill.fore_color.rgb = COLOR_WHITE
    hub_circle.line.color.rgb = COLOR_TEAL
    hub_circle.line.width = Pt(3)

    hub_tf = hub_circle.text_frame
    hub_tf.word_wrap = True
    p_h1 = hub_tf.paragraphs[0]
    p_h1.text = "🖥️\nDATACENTER"
    p_h1.font.size = Pt(17)
    p_h1.font.bold = True
    p_h1.font.color.rgb = COLOR_DARK
    p_h1.alignment = PP_ALIGN.CENTER

    p_h2 = hub_tf.add_paragraph()
    p_h2.text = "Cœur de Gestion des Données de Santé"
    p_h2.font.size = Pt(10.5)
    p_h2.font.color.rgb = COLOR_MUTED
    p_h2.alignment = PP_ALIGN.CENTER

    add_horizontal_pill(s3, Inches(5.4), Inches(1.5), Inches(6.7), Inches(1.15), "1", "Centraliser les données", "Entrepôt unifié national mettant fin aux silos d'informations.", COLOR_RED)
    add_horizontal_pill(s3, Inches(5.4), Inches(2.8), Inches(6.7), Inches(1.15), "2", "Sécuriser les informations", "Confidentialité, chiffrement, haute résilience et conformité SSI.", COLOR_TEAL)
    add_horizontal_pill(s3, Inches(5.4), Inches(4.1), Inches(6.7), Inches(1.15), "3", "Déployer des systèmes intégrés", "Interopérabilité fluide entre structures sanitaires et laboratoires.", COLOR_ORANGE)
    add_horizontal_pill(s3, Inches(5.4), Inches(5.4), Inches(6.7), Inches(1.15), "4", "Moderniser les outils de gestion", "Applications ergonomiques au service des praticiens et décideurs.", COLOR_GREEN)

    # ==================== SLIDE 4: Politique de Digitalisation ====================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, 4, "Politique de Digitalisation", "Vision Sectorielle")

    add_horizontal_pill(s4, Inches(0.8), Inches(1.4), Inches(3.6), Inches(1.15), "A", "Administration", "Pilotage RH, finances & gouvernance", COLOR_DARK)
    
    mid_badge = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.4), Inches(3.7), Inches(1.15))
    mid_badge.fill.solid()
    mid_badge.fill.fore_color.rgb = COLOR_WHITE
    mid_badge.line.color.rgb = COLOR_RED
    mid_badge.line.width = Pt(2)
    p_mb = mid_badge.text_frame.paragraphs[0]
    p_mb.text = "🌐  DIGITALISATION\nConvergence Stratégique"
    p_mb.font.size = Pt(11)
    p_mb.font.bold = True
    p_mb.font.color.rgb = COLOR_RED
    p_mb.alignment = PP_ALIGN.CENTER

    add_horizontal_pill(s4, Inches(8.9), Inches(1.4), Inches(3.6), Inches(1.15), "B", "Offre de soins", "Hôpitaux, cliniques & consultations", COLOR_TEAL)

    w4 = Inches(2.7)
    add_vertical_pill(s4, Inches(0.8), Inches(2.9), w4, Inches(3.6), "1", "Projets & Apps", "Déploiement des logiciels hospitaliers et applicatifs métiers.", COLOR_RED)
    add_vertical_pill(s4, Inches(3.8), Inches(2.9), w4, Inches(3.6), "2", "Sécurité SI", "Politique de sécurité, audits et protection des dossiers.", COLOR_TEAL)
    add_vertical_pill(s4, Inches(6.8), Inches(2.9), w4, Inches(3.6), "3", "Archivage", "Gestion électronique (GED) et traçabilité légale.", COLOR_ORANGE)
    add_vertical_pill(s4, Inches(9.8), Inches(2.9), w4, Inches(3.6), "4", "Données", "Consolidation nationale pour un pilotage continu.", COLOR_GREEN)

    # ==================== SLIDE 5: Transformation du SNIS ====================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, 5, "Transformation du SNIS", "Écosystème Applicatif")

    b5 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.4), Inches(10.9), Inches(0.95))
    b5.fill.solid()
    b5.fill.fore_color.rgb = COLOR_WHITE
    b5.line.color.rgb = COLOR_TEAL
    b5.line.width = Pt(2)
    p_b5 = b5.text_frame.paragraphs[0]
    p_b5.text = "🎯 Le SNIS évolue pour rendre les données : DISPONIBLES • FIABLES • ACCESSIBLES À TEMPS"
    p_b5.font.size = Pt(13)
    p_b5.font.bold = True
    p_b5.font.color.rgb = COLOR_DARK
    p_b5.alignment = PP_ALIGN.CENTER

    add_horizontal_pill(s5, Inches(1.2), Inches(2.6), Inches(10.9), Inches(1.25), "1", "DHIS2 — Collecte & Analyse Statistique", "Agrégation globale des données sanitaires, indicateurs nationaux et tableaux de bord de santé publique.", COLOR_RED)
    add_horizontal_pill(s5, Inches(1.2), Inches(4.05), Inches(10.9), Inches(1.25), "2", "SIGDEP3 — Suivi Clinique des Patients", "Système intégré pour le suivi médical et biologique individualisé des patients sous traitement ARV.", COLOR_TEAL)
    add_horizontal_pill(s5, Inches(1.2), Inches(5.5), Inches(10.9), Inches(1.25), "3", "OPENELIS — Gestion des Laboratoires", "Informatisation complète des laboratoires d'analyses médicales et traçabilité des flux d'échantillons.", COLOR_GREEN)

    # ==================== SLIDE 6: Autres Solutions Numériques ====================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, 6, "Autres Solutions Numériques", "Complémentarité Métier")

    add_horizontal_pill(s6, Inches(1.0), Inches(1.6), Inches(5.4), Inches(2.1), "1", "E-SIGL", "Gestion logistique intégrée des pharmacies et approvisionnement en intrants stratégiques.", COLOR_RED)
    add_horizontal_pill(s6, Inches(6.9), Inches(1.6), Inches(5.4), Inches(2.1), "2", "M-SUPPLY", "Gestion des stocks et traçabilité de la dispensation des médicaments au comptoir.", COLOR_TEAL)
    add_horizontal_pill(s6, Inches(1.0), Inches(4.0), Inches(5.4), Inches(2.1), "3", "UPID", "Identifiant Unique du Patient (notamment sous ARV) pour l'unicité et le suivi longitudinal.", COLOR_ORANGE)
    add_horizontal_pill(s6, Inches(6.9), Inches(4.0), Inches(5.4), Inches(2.1), "4", "DPI", "Dossier Patient Informatisé pour la continuité des soins et l'historique médical partagé.", COLOR_GREEN)

    b6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.5), Inches(6.35), Inches(8.333), Inches(0.5))
    b6.fill.solid()
    b6.fill.fore_color.rgb = COLOR_WHITE
    b6.line.color.rgb = COLOR_BORDER
    p_b6 = b6.text_frame.paragraphs[0]
    p_b6.text = "💡 Des solutions complémentaires au service de la transformation numérique."
    p_b6.font.size = Pt(11)
    p_b6.font.bold = True
    p_b6.font.color.rgb = COLOR_TEAL
    p_b6.alignment = PP_ALIGN.CENTER

    # ==================== SLIDE 7: Mise en Œuvre du Datacenter ====================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, 7, "Mise en Œuvre du Datacenter", "Socle d'Infrastructure")

    b7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.4), Inches(10.9), Inches(1.2))
    b7.fill.solid()
    b7.fill.fore_color.rgb = COLOR_WHITE
    b7.line.color.rgb = COLOR_TEAL
    b7.line.width = Pt(2)
    p_b7 = b7.text_frame.paragraphs[0]
    p_b7.text = "🖥️  INFRASTRUCTURE NATIONALE DE SANTÉ"
    p_b7.font.size = Pt(15)
    p_b7.font.bold = True
    p_b7.font.color.rgb = COLOR_DARK
    p_b7.alignment = PP_ALIGN.CENTER
    p_b7_sub = b7.text_frame.add_paragraph()
    p_b7_sub.text = "« Mettre en place une infrastructure centralisée, sécurisée et performante pour accompagner la transformation numérique du système de santé. »"
    p_b7_sub.font.size = Pt(11)
    p_b7_sub.font.color.rgb = COLOR_TEXT_DARK
    p_b7_sub.alignment = PP_ALIGN.CENTER

    w3 = Inches(3.4)
    add_vertical_pill(s7, Inches(1.2), Inches(2.9), w3, Inches(3.6), "1", "CENTRALISÉ", "Consolidation des données et mutualisation des ressources serveur.", COLOR_RED)
    add_vertical_pill(s7, Inches(4.95), Inches(2.9), w3, Inches(3.6), "2", "SÉCURISÉ", "Protection maximale, chiffrement et haute résilience contre les sinistres.", COLOR_TEAL)
    add_vertical_pill(s7, Inches(8.7), Inches(2.9), w3, Inches(3.6), "3", "PERFORMANT", "Temps de réponse optimaux et disponibilité continue 24/7.", COLOR_GREEN)

    # ==================== SLIDE 8: Chronologie de mise en œuvre ====================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, 8, "Chronologie de Mise en Œuvre (2021 - 2024)", "Feuille de Route")

    add_vertical_pill(s8, Inches(0.8), Inches(1.8), w4, Inches(4.6), "1", "2021 — Cadrage", "Signature des contrats, accords de financement et passation des marchés.", COLOR_RED, "Phase : Préparation")
    add_vertical_pill(s8, Inches(3.8), Inches(1.8), w4, Inches(4.6), "2", "2022 — Travaux", "Démarrage du génie civil, pose des baies serveurs et climatisation.", COLOR_TEAL, "Phase : Réalisation")
    add_vertical_pill(s8, Inches(6.8), Inches(1.8), w4, Inches(4.6), "3", "2023 — Recette", "Recette d'usine, tests de charge, bascule et validation sécuritaire.", COLOR_ORANGE, "Phase : Validation")
    add_vertical_pill(s8, Inches(9.8), Inches(1.8), w4, Inches(4.6), "4", "2024 — Service", "Réception provisoire et mise en service opérationnelle active.", COLOR_GREEN, "Phase : Exploitation")

    # ==================== SLIDE 9: Le Datacenter en Images (Using Real Photos!) ====================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, 9, "Le Datacenter en Images", "Galerie Visuelle")

    w_img = Inches(3.6)
    h_img = Inches(4.6)
    # Photo 1: datacenter1.png (Salle Serveurs)
    add_image_card(s9, Inches(0.8), Inches(1.8), w_img, h_img, img_dc1, "1", "Salle Serveurs & Baies IT", "Racks haute densité Schneider/APC", COLOR_RED)
    # Photo 2: datacenter4.png (Groupes électrogènes)
    add_image_card(s9, Inches(4.85), Inches(1.8), w_img, h_img, img_dc4, "2", "Énergie & Groupes Kohler", "Groupes électrogènes SDMO 165 kVA", COLOR_TEAL)
    # Photo 3: datacenter2.png (Monitoring)
    add_image_card(s9, Inches(8.9), Inches(1.8), w_img, h_img, img_dc2, "3", "Supervision & Monitoring", "Monitoring Schneider & Cogitech 24/7", COLOR_GREEN)

    # ==================== SLIDE 10: Infrastructure technique (Using Real Photos!) ====================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, 10, "Infrastructure Technique du Datacenter", "Spécifications Équipements")

    # Photo 1: datacenter1.png (Puissance IT)
    add_image_card(s10, Inches(0.8), Inches(1.8), w_img, h_img, img_dc1, "1", "Baies Serveurs & Câblage", "Confinement InRow & serveurs SNIS", COLOR_RED)
    # Photo 2: datacenter3.png (Climatisation externe)
    add_image_card(s10, Inches(4.85), Inches(1.8), w_img, h_img, img_dc3, "2", "Aéroréfrigérants Précis", "Refroidissement redondant dédié", COLOR_TEAL)
    # Photo 3: datacenter4.png (Secours énergie)
    add_image_card(s10, Inches(8.9), Inches(1.8), w_img, h_img, img_dc4, "3", "Continuité Électrique", "Automates de bascule & secours", COLOR_GREEN)

    # ==================== SLIDE 11: Une infrastructure au service de la santé ====================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, 11, "Au Service de la Santé Digitale", "Valeur Ajoutée")

    # Left Image Card with datacenter_slide_11_image_3.png
    card_l11 = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.2), Inches(4.8))
    card_l11.fill.solid()
    card_l11.fill.fore_color.rgb = COLOR_WHITE
    card_l11.line.color.rgb = COLOR_BORDER
    card_l11.line.width = Pt(1.5)

    if os.path.exists(img_dc11):
        s11.shapes.add_picture(img_dc11, Inches(0.95), Inches(1.75), width=Inches(4.9), height=Inches(3.3))
    elif os.path.exists(img_dc1):
        s11.shapes.add_picture(img_dc1, Inches(0.95), Inches(1.75), width=Inches(4.9), height=Inches(3.3))

    tx_l11 = s11.shapes.add_textbox(Inches(0.95), Inches(5.2), Inches(4.9), Inches(1.0))
    p_l11 = tx_l11.text_frame.paragraphs[0]
    p_l11.text = "🛡️  SOCLE DE CONFIANCE & HAUTE SÉCURITÉ"
    p_l11.font.size = Pt(11)
    p_l11.font.bold = True
    p_l11.font.color.rgb = COLOR_DARK
    p_l11.alignment = PP_ALIGN.CENTER
    p_l11_sub = tx_l11.text_frame.add_paragraph()
    p_l11_sub.text = "Salle serveurs climatisée, extinction gaz inerte et contrôle d'accès sécurisé."
    p_l11_sub.font.size = Pt(9.5)
    p_l11_sub.font.color.rgb = COLOR_MUTED
    p_l11_sub.alignment = PP_ALIGN.CENTER
    
    add_horizontal_pill(s11, Inches(6.4), Inches(1.6), Inches(6.1), Inches(1.4), "1", "Centralisation des Données", "Unification complète des bases de données sanitaires nationales.", COLOR_RED)
    add_horizontal_pill(s11, Inches(6.4), Inches(3.3), Inches(6.1), Inches(1.4), "2", "Sécurisation & Résilience", "Chiffrement et conformité stricte aux exigences de cybersécurité.", COLOR_TEAL)
    add_horizontal_pill(s11, Inches(6.4), Inches(5.0), Inches(6.1), Inches(1.4), "3", "Performance & Disponibilité", "Temps d'accès instantanés pour les professionnels de santé.", COLOR_GREEN)

    # ==================== SLIDE 12: Défis et Perspectives ====================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, 12, "Défis et Perspectives", "Enjeux Stratégiques")

    w5 = Inches(2.25)
    add_vertical_pill(s12, Inches(0.6), Inches(1.8), w5, Inches(4.6), "1", "Ressources Humaines", "Recrutement et fidélisation d'experts IT qualifiés.", COLOR_RED)
    add_vertical_pill(s12, Inches(3.05), Inches(1.8), w5, Inches(4.6), "2", "Formation Continue", "Montée en compétences des agents et soignants.", COLOR_TEAL)
    add_vertical_pill(s12, Inches(5.5), Inches(1.8), w5, Inches(4.6), "3", "Maintenance", "Maintenance préventive, SLA et pièces de rechange.", COLOR_ORANGE)
    add_vertical_pill(s12, Inches(7.95), Inches(1.8), w5, Inches(4.6), "4", "Interconnexion", "Raccordement réseau des centres périphériques.", COLOR_GREEN)
    add_vertical_pill(s12, Inches(10.4), Inches(1.8), w5, Inches(4.6), "5", "Sécurité Données", "Conformité PSSI, audits et cyber-résilience continue.", COLOR_DARK)

    # ==================== SLIDE 13: Connectivité et Interconnexion ====================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, 13, "Connectivité & Interconnexion", "Maillage Réseau")

    w6 = Inches(5.5)
    h6 = Inches(1.4)
    add_horizontal_pill(s13, Inches(0.8), Inches(1.6), w6, h6, "1", "Fibre Optique", "Dorsale nationale haut débit pour les CHU & CHR.", COLOR_RED)
    add_horizontal_pill(s13, Inches(6.9), Inches(1.6), w6, h6, "2", "VSAT Gouvernemental", "Liaisons satellitaires pour les localités enclavées.", COLOR_TEAL)
    add_horizontal_pill(s13, Inches(0.8), Inches(3.3), w6, h6, "3", "Internet Haut Débit", "Garantie de bande passante et latence minimale.", COLOR_ORANGE)
    add_horizontal_pill(s13, Inches(6.9), Inches(3.3), w6, h6, "4", "Couverture 4G", "Tablettes terrain et accès mobile pour la collecte.", COLOR_GREEN)
    add_horizontal_pill(s13, Inches(0.8), Inches(5.0), w6, h6, "5", "Coût de la Data", "Tarification préférentielle et partenariats télécoms.", COLOR_DARK)
    add_horizontal_pill(s13, Inches(6.9), Inches(5.0), w6, h6, "6", "Alimentation Électrique", "Stabilité du réseau et solutions solaires d'appoint.", COLOR_RED)

    # ==================== SLIDE 14: Identité, Adoption et Sécurité ====================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, 14, "Identité, Adoption et Sécurité", "Gouvernance & Confiance")

    add_vertical_pill(s14, Inches(0.8), Inches(1.8), w4, Inches(4.6), "1", "Identifiant Unique", "Garantir l'unicité du patient et éviter les doublons.", COLOR_RED, "→ IDENTIFIER")
    add_vertical_pill(s14, Inches(3.8), Inches(1.8), w4, Inches(4.6), "2", "Appropriation Outils", "Conduite du changement active et accompagnement.", COLOR_TEAL, "→ ADOPTER")
    add_vertical_pill(s14, Inches(6.8), Inches(1.8), w4, Inches(4.6), "3", "Sécurité des Données", "Protection du secret médical, chiffrement & intégrité.", COLOR_ORANGE, "→ SÉCURISER")
    add_vertical_pill(s14, Inches(9.8), Inches(1.8), w4, Inches(4.6), "4", "Cadre Réglementaire", "Conformité légale aux lois sur les données (ARTCI).", COLOR_GREEN, "→ ENCADRER")

    # ==================== SLIDE 15: Conclusion ====================
    s15 = prs.slides.add_slide(blank_layout)
    add_header(s15, 15, "Conclusion & Perspectives", "Clôture")

    b15 = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.4), Inches(10.9), Inches(1.4))
    b15.fill.solid()
    b15.fill.fore_color.rgb = COLOR_WHITE
    b15.line.color.rgb = COLOR_TEAL
    b15.line.width = Pt(2)
    p15 = b15.text_frame.paragraphs[0]
    p15.text = "« La digitalisation transforme progressivement le système de santé ivoirien. »"
    p15.font.size = Pt(21)
    p15.font.bold = True
    p15.font.color.rgb = COLOR_DARK
    p15.alignment = PP_ALIGN.CENTER

    add_vertical_pill(s15, Inches(0.8), Inches(3.0), w4, Inches(2.8), "1", "Accès à l'Information", "Données sanitaires fiables en temps réel pour tous les acteurs.", COLOR_RED)
    add_vertical_pill(s15, Inches(3.8), Inches(3.0), w4, Inches(2.8), "2", "Qualité des Soins", "Prise en charge patient fluidifiée et suivi médical continu.", COLOR_TEAL)
    add_vertical_pill(s15, Inches(6.8), Inches(3.0), w4, Inches(2.8), "3", "Gestion du Système", "Efficience accrue des ressources matérielles et humaines.", COLOR_ORANGE)
    add_vertical_pill(s15, Inches(9.8), Inches(3.0), w4, Inches(2.8), "4", "Prise de Décision", "Politiques publiques fondées sur des preuves tangibles.", COLOR_GREEN)

    cta = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(6.1), Inches(10.333), Inches(0.6))
    cta.fill.solid()
    cta.fill.fore_color.rgb = COLOR_WHITE
    cta.line.color.rgb = COLOR_GREEN
    p_cta = cta.text_frame.paragraphs[0]
    p_cta.text = "✨ La réussite du projet repose sur l'engagement et la contribution de tous les acteurs."
    p_cta.font.size = Pt(12)
    p_cta.font.bold = True
    p_cta.font.color.rgb = COLOR_GREEN
    p_cta.alignment = PP_ALIGN.CENTER

    output_path = "Presentation_Datacenter_Sante_CI_Style.pptx"
    prs.save(output_path)
    print(f"Presentation successfully updated with real images and saved to {output_path}")

if __name__ == "__main__":
    create_styled_presentation()
